#!/usr/bin/env python3
"""Regenerate the test documents (developer tool; the outputs are committed).

    pip install python-docx reportlab openpyxl pillow
    python make_fixtures.py

The documents are invented ("ACME EM-100" does not exist) but laid out the way
real manuals are: register notation, hex addresses with ranges, read-only and
write-only rows, reserved rows, a function-code caption above each table, and
columns that are sometimes empty.
"""
import csv
import io

T1 = [
    ("Register", "Description", "Type", "Unit"),
    ("30001", "Voltage L1", "Float32", "V"),
    ("30003", "Voltage L2", "Float32", "V"),
    ("30005", "Voltage L3", "Float32", "V"),
    ("30007", "Total active power", "Float32", "W"),
    ("30009", "Frequency", "Float32", "Hz"),
]
T2 = [
    ("Reg (hex)", "Parameter", "Data type", "Multiplier", "R/W"),
    ("0x0000", "Voltage L1", "UINT16", "0.1", "R"),
    ("0x0001", "Current L1", "UINT16", "0.01", "R"),
    ("0x0002", "Active power", "INT16", "", "R"),
    ("0x0003", "Temperature (°C)", "INT16", "0.1", "R"),
    ("0x0004-0x0005", "Energy import", "UINT32", "0.01", "R"),
    ("0x0010", "Reset energy counters", "UINT16", "", "W"),
    ("0x0011", "Reserved", "", "", "R"),
]
TR = [
    ("Adres", "Parametre", "Birim", "Veri Tipi", "Açıklama"),
    ("40101", "Şebeke gerilimi L1", "V", "Float", "Faz-nötr"),
    ("40103", "Şebeke gerilimi L2", "V", "Float", "Faz-nötr"),
    ("40105", "Frekans", "Hz", "Float", ""),
    ("40107", "Güç faktörü", "", "Float", ""),
]


def make_docx():
    from docx import Document
    d = Document()
    d.add_heading("ACME EM-100 Modbus register map", 1)
    d.add_paragraph("32-bit values use big-endian word order (ABCD).")
    d.add_paragraph("Table 1 - Input registers (function code 04)")
    for rows in (T1,):
        t = d.add_table(rows=len(rows), cols=len(rows[0]))
        for i, r in enumerate(rows):
            for j, c in enumerate(r):
                t.cell(i, j).text = c
    d.add_paragraph("Table 2 - Holding registers (function code 03)")
    t = d.add_table(rows=len(T2), cols=len(T2[0]))
    for i, r in enumerate(T2):
        for j, c in enumerate(r):
            t.cell(i, j).text = c
    d.save("acme-em100.docx")

    d = Document()
    d.add_paragraph("DKM Modbus adres tablosu - Tutma register (fonksiyon 03), Float, ABCD")
    t = d.add_table(rows=len(TR), cols=len(TR[0]))
    for i, r in enumerate(TR):
        for j, c in enumerate(r):
            t.cell(i, j).text = c
    d.save("datakom-tr.docx")


def make_pdf():
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    c = canvas.Canvas("acme-em100.pdf", pagesize=A4)
    w, h = A4
    y = h - 60
    c.setFont("Helvetica-Bold", 14)
    c.drawString(50, y, "ACME EM-100 Modbus register map")
    c.setFont("Helvetica", 9)
    y -= 24
    c.drawString(50, y, "32-bit values use big-endian word order (ABCD).")
    y -= 24
    c.drawString(50, y, "Table 1 - Input registers (function code 04)")
    y -= 18
    xs1 = (50, 130, 360, 450)
    for ri, row in enumerate(T1):
        if ri == 0:
            c.setFont("Helvetica-Bold", 9)
        else:
            c.setFont("Helvetica", 9)
        for x, cell in zip(xs1, row):
            if cell:
                c.drawString(x, y, cell)
        y -= 14
    y -= 14
    c.setFont("Helvetica", 9)
    c.drawString(50, y, "Table 2 - Holding registers (function code 03)")
    y -= 18
    xs2 = (50, 130, 300, 380, 470)
    for ri, row in enumerate(T2):
        c.setFont("Helvetica-Bold" if ri == 0 else "Helvetica", 9)
        for x, cell in zip(xs2, row):
            if cell:                       # empty cells are simply not drawn, as in real PDFs
                c.drawString(x, y, cell)
        y -= 14
    c.showPage()
    c.save()

    # A page with no text at all: stands in for a scanned manual.
    c = canvas.Canvas("scanned.pdf", pagesize=A4)
    c.rect(100, 400, 200, 100, stroke=1, fill=0)
    c.showPage()
    c.save()


def make_xlsx():
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Registers"
    ws.append(["ACME EM-100 holding registers (function code 03)"])
    ws.append([])
    for r in T2:
        ws.append([int(r[0], 16) if r[0].startswith("0x") and "-" not in r[0] else r[0], *r[1:]])
    wb.create_sheet("Notes").append(["nothing useful here"])
    wb.save("acme-em100.xlsx")


def make_csv():
    out = io.StringIO()
    w = csv.writer(out, delimiter=";", lineterminator="\n")
    w.writerow(["Adres", "Parametre", "Tip", "Birim", "Çarpan"])
    w.writerow(["40001", "Gerilim L1", "UINT16", "V", "0,1"])
    w.writerow(["40002", "Akım L1", "UINT16", "A", "0,01"])
    w.writerow(["40003", "Aktif güç", "INT16", "W", "1"])
    w.writerow(["40004", "Enerji", "UINT32", "kWh", "x100"])
    open("datakom-tr.csv", "w", encoding="utf-8", newline="").write(out.getvalue())


if __name__ == "__main__":
    make_docx()
    make_pdf()
    make_xlsx()
    make_csv()
    print("fixtures written")
